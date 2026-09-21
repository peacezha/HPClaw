from __future__ import annotations

import re
import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "papers" / "hpclaw_paper_draft.md"
OUTPUT = ROOT / "docs" / "papers" / "hpclaw_paper_draft.docx"

BODY_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"
INK = RGBColor(35, 35, 35)
BLUE = RGBColor(35, 35, 35)
DARK_BLUE = RGBColor(35, 35, 35)
MUTED = RGBColor(95, 95, 95)
CALLOUT_FILL = "F4F6F9"
CALLOUT_BORDER = "D6DEE8"


def set_run_font(run, font_name: str = BODY_FONT, latin_name: str = LATIN_FONT, size: int | None = None) -> None:
    run.font.name = latin_name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin_name)
    rfonts.set(qn("w:hAnsi"), latin_name)
    rfonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 6, line: int = 280) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line / 240


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = CALLOUT_BORDER, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_inline_runs(paragraph, text: str, size: int | None = None, color: RGBColor | None = None) -> None:
    # Minimal Markdown inline support for backtick code and bold spans.
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=size)
            if color:
                run.font.color.rgb = color
        token = match.group(0)
        clean = token[1:-1] if token.startswith("`") else token[2:-2]
        run = paragraph.add_run(clean)
        set_run_font(run, size=size)
        if token.startswith("`"):
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt((size or 11) - 0.5)
            run.font.color.rgb = RGBColor(150, 70, 40)
        else:
            run.bold = True
            if color:
                run.font.color.rgb = color
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size)
        if color:
            run.font.color.rgb = color


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.allow_autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=2, after=2, line=280)
    add_inline_runs(p, text, size=10, color=MUTED)


def split_markdown_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_markdown_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def looks_like_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows = [split_markdown_table_row(line) for line in table_lines if not is_table_separator(line)]
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.autofit = False
    table.allow_autofit = False
    try:
        table.style = "Table Grid"
    except Exception:
        pass

    col_width = Inches(6.5 / max(1, col_count))
    for row_idx, row in enumerate(rows):
        for col_idx in range(col_count):
            cell = table.cell(row_idx, col_idx)
            cell.width = col_width
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, "F4F6F9")
            set_cell_border(cell, color="D6DEE8", size="6")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, before=0, after=0, line=260)
            text = row[col_idx] if col_idx < len(row) else ""
            add_inline_runs(p, text, size=8 if col_count >= 5 else 9, color=INK)
            if row_idx == 0:
                for run in p.runs:
                    run.bold = True

    after = doc.add_paragraph()
    set_paragraph_spacing(after, before=2, after=6, line=260)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title_block(doc: Document, lines: list[str]) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, before=0, after=10, line=300)
    title_run = title.add_run(lines[0].lstrip("# ").strip())
    set_run_font(title_run, size=20)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(20, 32, 52)

    for line in lines[1:]:
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, before=0, after=2, line=260)
        run = p.add_run(line.strip())
        set_run_font(run, size=10)
        run.font.color.rgb = MUTED


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HPClaw manuscript")
    set_run_font(run, size=9)
    run.font.color.rgb = MUTED


def add_list_item(doc: Document, text: str, ordered: bool) -> None:
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_paragraph_spacing(p, before=0, after=6, line=280)
    add_inline_runs(p, text)


def build_docx(source: Path = SOURCE, output: Path = OUTPUT) -> None:
    source_text = source.read_text(encoding="utf-8")
    lines = source_text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    first_break = 0
    for i, line in enumerate(lines):
        if line.startswith("## "):
            first_break = i
            break
    add_title_block(doc, lines[:first_break])
    add_footer(doc)

    i = first_break
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        i += 1
        if not line:
            continue
        if looks_like_table_line(line):
            table_lines = [line]
            while i < len(lines) and looks_like_table_line(lines[i]):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=3)
        elif re.match(r"^\d+\.\s+", line):
            add_list_item(doc, re.sub(r"^\d+\.\s+", "", line), ordered=True)
            continue
        elif line.startswith("- "):
            add_list_item(doc, line[2:], ordered=False)
            continue
        elif line.startswith("> "):
            add_callout(doc, line[2:].strip())
            continue
        elif re.match(r"^图\s+\d+\s+生成提示词", line):
            add_callout(doc, line)
            continue
        elif re.match(r"^图\s+\d+\s+待补充", line):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=8, after=4, line=280)
            run = p.add_run(line)
            set_run_font(run, size=10)
            run.bold = True
            run.font.color.rgb = DARK_BLUE
            continue
        else:
            p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=6, line=280)
        if not p.runs:
            add_inline_runs(p, line)

    # Avoid a dangling final blank paragraph after table/callout-heavy content.
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1]._element.getparent().remove(doc.paragraphs[-1]._element)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert an HPClaw markdown paper draft to DOCX.")
    parser.add_argument("source", nargs="?", type=Path, default=SOURCE)
    parser.add_argument("output", nargs="?", type=Path, default=OUTPUT)
    args = parser.parse_args()
    build_docx(args.source, args.output)
    print(args.output)
