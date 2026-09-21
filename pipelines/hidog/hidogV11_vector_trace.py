#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Independent sgRNA/vector tracing workflow for HiDOG V11."""

import argparse
import csv
import gzip
import json
import os
import re
from collections import Counter
from dataclasses import asdict, dataclass
from typing import Dict, List, Optional, Sequence, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


DNA_BASES = frozenset("ACGT")
MANIFEST_FIELDS = (
    "vector_id",
    "guide_id",
    "spacer_seq",
    "gene_name",
    "reference_group",
)


@dataclass(frozen=True)
class GuideRecord:
    vector_id: str
    guide_id: str
    spacer_seq: str
    gene_name: str
    reference_group: str


@dataclass(frozen=True)
class SpacerRecord:
    spacer_id: str
    spacer_seq: str


@dataclass(frozen=True)
class SpacerMatch:
    match_type: str
    guide_id: Optional[str] = None
    orientation: Optional[str] = None
    candidate_guide_ids: Tuple[str, ...] = ()


@dataclass(frozen=True)
class VectorTraceConfig:
    anchor: str
    spacer_length: int = 20
    spacer_side: str = "after"
    anchor_read: str = "r1"
    anchor_max_mismatches: int = 0
    barcode_spacer_length: int = 4
    barcode_length: int = 4
    barcode_q30: int = 30
    min_guide_fraction: float = 0.05
    min_guide_reads: int = 10
    min_sample_barcode_reads: int = 2000
    min_sample_anchor_reads: int = 4000
    allow_unique_1mm: bool = True
    allow_reverse_complement: bool = False
    top_unmatched: int = 20

    def __post_init__(self):
        anchor = self.anchor.strip().upper()
        if not anchor or set(anchor) - DNA_BASES:
            raise ValueError("anchor must contain only A/C/G/T")
        if self.spacer_length <= 0:
            raise ValueError("spacer_length must be positive")
        if self.spacer_side not in {"before", "after"}:
            raise ValueError("spacer_side must be 'before' or 'after'")
        if self.anchor_read not in {"r1", "r2"}:
            raise ValueError("anchor_read must be 'r1' or 'r2'")
        if self.anchor_max_mismatches < 0:
            raise ValueError("anchor_max_mismatches cannot be negative")
        if self.barcode_spacer_length < 0 or self.barcode_length <= 0:
            raise ValueError("barcode layout lengths are invalid")
        if not 0 <= self.barcode_q30 <= 93:
            raise ValueError("barcode_q30 must be between 0 and 93")
        if not 0 < self.min_guide_fraction <= 1:
            raise ValueError("min_guide_fraction must be greater than 0 and at most 1")
        if self.min_guide_reads <= 0:
            raise ValueError("min_guide_reads must be positive")
        if self.min_sample_barcode_reads <= 0:
            raise ValueError("min_sample_barcode_reads must be positive")
        if self.min_sample_anchor_reads <= 0:
            raise ValueError("min_sample_anchor_reads must be positive")
        if self.top_unmatched < 0:
            raise ValueError("top_unmatched cannot be negative")
        object.__setattr__(self, "anchor", anchor)


@dataclass
class VectorTraceResult:
    sample_names: List[str]
    guides: List[GuideRecord]
    spacer_rows: List[dict]
    sample_rows: List[dict]
    sample_qc_rows: List[dict]
    sample_detail_rows: Dict[str, List[dict]]
    run_metrics: Dict[str, object]


def reverse_complement(sequence: str) -> str:
    return sequence.upper().translate(str.maketrans("ACGT", "TGCA"))[::-1]


def hamming_distance(left: str, right: str) -> int:
    if len(left) != len(right):
        raise ValueError("Hamming distance requires sequences with equal length")
    return sum(a != b for a, b in zip(left, right))


def _read_spacer_reference(path: str):
    with open(path, "r", encoding="utf-8-sig") as handle:
        lines = [(line_no, raw.strip()) for line_no, raw in enumerate(handle, start=1)]
    nonempty = [(line_no, line) for line_no, line in lines if line]
    if not nonempty:
        raise ValueError(f"spacer reference contains no records: {path}")

    if any(line.startswith(">") for _, line in nonempty):
        records = []
        header = None
        sequence_parts = []
        for line_no, line in nonempty:
            if line.startswith(">"):
                if header is not None:
                    if not sequence_parts:
                        raise ValueError(
                            f"spacer reference record has empty sequence: {header}"
                        )
                    records.append((header, "".join(sequence_parts)))
                header = line[1:].strip()
                if not header:
                    raise ValueError(
                        f"spacer reference line {line_no} has an empty FASTA header"
                    )
                sequence_parts = []
            else:
                if header is None:
                    raise ValueError(
                        f"spacer reference line {line_no} precedes the first FASTA header"
                    )
                sequence_parts.append(line)
        if header is not None:
            if not sequence_parts:
                raise ValueError(f"spacer reference record has empty sequence: {header}")
            records.append((header, "".join(sequence_parts)))
        return records

    basename = os.path.basename(path)
    return [(f"{basename}:{line_no}", line) for line_no, line in nonempty]


def load_spacer_references(
    paths: Sequence[str],
    spacer_length: int = 20,
) -> List[SpacerRecord]:
    spacers = []
    seen_ids = set()
    seen_sequences = set()
    for path in paths:
        for spacer_id, raw_sequence in _read_spacer_reference(path):
            sequence = raw_sequence.replace("U", "T").upper()
            if spacer_id in seen_ids:
                raise ValueError(f"duplicate spacer_id: {spacer_id}")
            if len(sequence) != spacer_length:
                raise ValueError(
                    f"spacer {spacer_id} must be {spacer_length} nt"
                )
            if set(sequence) - DNA_BASES:
                raise ValueError(f"spacer {spacer_id} contains non-ACGT bases")
            if sequence in seen_sequences:
                raise ValueError(f"duplicate spacer_seq: {sequence}")
            spacers.append(SpacerRecord(spacer_id=spacer_id, spacer_seq=sequence))
            seen_ids.add(spacer_id)
            seen_sequences.add(sequence)
    if not spacers:
        raise ValueError("spacer references contain no records")
    return spacers


def load_guide_manifest(path: str, spacer_length: int = 20) -> List[GuideRecord]:
    guides = []
    seen_guides = set()
    seen_spacers = set()
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        missing = [field for field in MANIFEST_FIELDS if field not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(
                "guide manifest is missing required columns: " + ", ".join(missing)
            )
        for line_no, row in enumerate(reader, start=2):
            values = {field: (row.get(field) or "").strip() for field in MANIFEST_FIELDS}
            empty = [field for field, value in values.items() if not value]
            if empty:
                raise ValueError(
                    f"guide manifest line {line_no} has empty fields: {', '.join(empty)}"
                )

            spacer = values["spacer_seq"].upper()
            if len(spacer) != spacer_length:
                raise ValueError(
                    f"guide manifest line {line_no} spacer_seq must be {spacer_length} nt"
                )
            if set(spacer) - DNA_BASES:
                raise ValueError(
                    f"guide manifest line {line_no} spacer_seq contains non-ACGT bases"
                )
            if values["guide_id"] in seen_guides:
                raise ValueError(
                    f"guide manifest line {line_no} has duplicate guide_id: {values['guide_id']}"
                )
            if spacer in seen_spacers:
                raise ValueError(
                    f"guide manifest line {line_no} has duplicate spacer_seq: {spacer}"
                )

            guide = GuideRecord(
                vector_id=values["vector_id"],
                guide_id=values["guide_id"],
                spacer_seq=spacer,
                gene_name=values["gene_name"],
                reference_group=values["reference_group"],
            )
            guides.append(guide)
            seen_guides.add(guide.guide_id)
            seen_spacers.add(guide.spacer_seq)

    if not guides:
        raise ValueError("guide manifest contains no guide records")
    return guides


def _matching_orientations(
    sequence: str,
    guide: GuideRecord,
    allow_reverse_complement: bool,
):
    yield guide.spacer_seq, "DIRECT"
    if allow_reverse_complement:
        reverse = reverse_complement(guide.spacer_seq)
        if reverse != guide.spacer_seq:
            yield reverse, "REVERSE_COMPLEMENT"


def _resolve_match(candidates, match_type: str) -> SpacerMatch:
    by_guide = {}
    for guide, orientation in candidates:
        if guide.guide_id not in by_guide or orientation == "DIRECT":
            by_guide[guide.guide_id] = (guide, orientation)
    if len(by_guide) == 1:
        guide, orientation = next(iter(by_guide.values()))
        return SpacerMatch(match_type, guide.guide_id, orientation, (guide.guide_id,))
    if len(by_guide) > 1:
        return SpacerMatch(
            "AMBIGUOUS",
            candidate_guide_ids=tuple(sorted(by_guide)),
        )
    return SpacerMatch("UNMATCHED")


def classify_spacer(
    sequence: str,
    guides: Sequence[GuideRecord],
    allow_unique_1mm: bool = True,
    allow_reverse_complement: bool = False,
) -> SpacerMatch:
    sequence = sequence.strip().upper()
    if not sequence or set(sequence) - DNA_BASES:
        return SpacerMatch("UNMATCHED")

    exact_candidates = []
    for guide in guides:
        for target, orientation in _matching_orientations(
            sequence,
            guide,
            allow_reverse_complement,
        ):
            if sequence == target:
                exact_candidates.append((guide, orientation))
    exact = _resolve_match(exact_candidates, "EXACT")
    if exact.match_type != "UNMATCHED":
        return exact
    if not allow_unique_1mm:
        return SpacerMatch("UNMATCHED")

    mismatch_candidates = []
    for guide in guides:
        for target, orientation in _matching_orientations(
            sequence,
            guide,
            allow_reverse_complement,
        ):
            if len(sequence) == len(target) and hamming_distance(sequence, target) == 1:
                mismatch_candidates.append((guide, orientation))
    return _resolve_match(mismatch_candidates, "UNIQUE_1MM")


def extract_spacer_from_read(
    sequence: str,
    anchor: str,
    spacer_length: int = 20,
    spacer_side: str = "after",
    anchor_max_mismatches: int = 0,
):
    sequence = sequence.strip().upper()
    anchor = anchor.strip().upper()
    if not anchor or set(anchor) - DNA_BASES:
        raise ValueError("anchor must contain only A/C/G/T")
    if spacer_length <= 0:
        raise ValueError("spacer_length must be positive")
    if spacer_side not in {"before", "after"}:
        raise ValueError("spacer_side must be 'before' or 'after'")
    if anchor_max_mismatches < 0:
        raise ValueError("anchor_max_mismatches cannot be negative")

    candidates = []
    for start in range(0, len(sequence) - len(anchor) + 1):
        observed = sequence[start:start + len(anchor)]
        mismatches = hamming_distance(observed, anchor)
        if mismatches <= anchor_max_mismatches:
            candidates.append((mismatches, start))
    if not candidates:
        return None, "ANCHOR_NOT_FOUND"

    best_mismatches = min(item[0] for item in candidates)
    best_positions = [start for mismatches, start in candidates if mismatches == best_mismatches]
    if len(best_positions) != 1:
        return None, "ANCHOR_AMBIGUOUS"

    anchor_start = best_positions[0]
    if spacer_side == "after":
        spacer_start = anchor_start + len(anchor)
        spacer_end = spacer_start + spacer_length
    else:
        spacer_end = anchor_start
        spacer_start = spacer_end - spacer_length
    if spacer_start < 0 or spacer_end > len(sequence):
        return None, "SPACER_LENGTH_INVALID"

    spacer = sequence[spacer_start:spacer_end]
    if set(spacer) - DNA_BASES:
        return None, "SPACER_INVALID_BASES"
    return spacer, "EXTRACTED"


def _open_fastq_text(path):
    if path.lower().endswith(".gz"):
        return gzip.open(path, "rt", encoding="utf-8")
    return open(path, "r", encoding="utf-8")


def _iter_fastq(path):
    with _open_fastq_text(path) as handle:
        record_no = 0
        while True:
            name = handle.readline()
            if not name:
                break
            sequence = handle.readline()
            plus = handle.readline()
            quality = handle.readline()
            record_no += 1
            if not sequence or not plus or not quality:
                raise ValueError(f"truncated FASTQ record {record_no}: {path}")
            if not name.startswith("@") or not plus.startswith("+"):
                raise ValueError(f"invalid FASTQ record {record_no}: {path}")
            sequence = sequence.rstrip("\r\n").upper()
            quality = quality.rstrip("\r\n")
            if len(sequence) != len(quality):
                raise ValueError(f"FASTQ sequence/quality length mismatch at record {record_no}: {path}")
            yield name.rstrip("\r\n"), sequence, quality


def _normalized_fastq_name(name):
    value = name[1:] if name.startswith("@") else name
    value = value.split()[0]
    if value.endswith("/1") or value.endswith("/2"):
        value = value[:-2]
    return value


def _barcode_region_passes(quality, start, length, threshold):
    end = start + length
    if len(quality) < end:
        return False
    return all(ord(char) - 33 >= threshold for char in quality[start:end])


def load_vector_barcodes(path, barcode_length=4):
    barcode_map = {}
    sample_names = []
    seen_samples = set()
    with open(path, "r", encoding="utf-8-sig") as handle:
        for line_no, raw_line in enumerate(handle, start=1):
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            fields = line.split()
            if len(fields) < 3:
                raise ValueError(
                    f"barcode file line {line_no} requires: sample_name barcode_R1 barcode_R2"
                )
            sample_id, barcode_r1, barcode_r2 = fields[:3]
            barcode_r1 = barcode_r1.upper()
            barcode_r2 = barcode_r2.upper()
            if sample_id in seen_samples:
                raise ValueError(f"duplicate sample_id in barcode file: {sample_id}")
            if len(barcode_r1) != barcode_length or len(barcode_r2) != barcode_length:
                raise ValueError(
                    f"barcode file line {line_no} barcodes must be {barcode_length} nt"
                )
            if set(barcode_r1 + barcode_r2) - DNA_BASES:
                raise ValueError(f"barcode file line {line_no} contains non-ACGT bases")
            key = barcode_r1 + barcode_r2
            if key in barcode_map:
                raise ValueError(f"duplicate barcode pair at line {line_no}")
            barcode_map[key] = sample_id
            sample_names.append(sample_id)
            seen_samples.add(sample_id)
    if not sample_names:
        raise ValueError("barcode file contains no sample records")
    return barcode_map, sample_names


def analyze_vector_trace(
    read1,
    read2,
    barcode_file,
    guides,
    config,
):
    barcode_map, sample_names = load_vector_barcodes(
        barcode_file,
        barcode_length=config.barcode_length,
    )
    extracted_sequences = {sample: Counter() for sample in sample_names}
    sample_metrics = {sample: Counter() for sample in sample_names}
    run_metrics = Counter()
    barcode_start = config.barcode_spacer_length

    r1_iter = _iter_fastq(read1)
    r2_iter = _iter_fastq(read2)
    while True:
        try:
            r1_record = next(r1_iter)
        except StopIteration:
            r1_record = None
        try:
            r2_record = next(r2_iter)
        except StopIteration:
            r2_record = None
        if r1_record is None and r2_record is None:
            break
        if r1_record is None or r2_record is None:
            raise ValueError("R1 and R2 FASTQ files are not synchronized")

        r1_name, r1_sequence, r1_quality = r1_record
        r2_name, r2_sequence, r2_quality = r2_record
        run_metrics["total_pairs"] += 1
        if _normalized_fastq_name(r1_name) != _normalized_fastq_name(r2_name):
            raise ValueError(
                f"R1 and R2 FASTQ names are not synchronized: {r1_name} vs {r2_name}"
            )
        if not _barcode_region_passes(
            r1_quality,
            barcode_start,
            config.barcode_length,
            config.barcode_q30,
        ):
            run_metrics["barcode_q30_fail_r1"] += 1
            continue
        if not _barcode_region_passes(
            r2_quality,
            barcode_start,
            config.barcode_length,
            config.barcode_q30,
        ):
            run_metrics["barcode_q30_fail_r2"] += 1
            continue

        barcode_end = barcode_start + config.barcode_length
        barcode_key = r1_sequence[barcode_start:barcode_end] + r2_sequence[barcode_start:barcode_end]
        sample_id = barcode_map.get(barcode_key)
        if sample_id is None:
            run_metrics["barcode_mismatch"] += 1
            continue
        run_metrics["barcode_matched_pairs"] += 1
        sample_metrics[sample_id]["barcode_matched_pairs"] += 1

        anchor_sequence = r1_sequence if config.anchor_read == "r1" else r2_sequence
        spacer, extraction_status = extract_spacer_from_read(
            anchor_sequence,
            anchor=config.anchor,
            spacer_length=config.spacer_length,
            spacer_side=config.spacer_side,
            anchor_max_mismatches=config.anchor_max_mismatches,
        )
        run_metrics[extraction_status.lower()] += 1
        sample_metrics[sample_id][extraction_status.lower()] += 1
        if extraction_status != "EXTRACTED":
            continue
        extracted_sequences[sample_id][spacer] += 1

    exact_counts = {sample: Counter() for sample in sample_names}
    rescued_counts = {sample: Counter() for sample in sample_names}
    sequence_matches = {sample: {} for sample in sample_names}
    for sample_id in sample_names:
        for spacer, count in extracted_sequences[sample_id].items():
            match = classify_spacer(
                spacer,
                guides,
                allow_unique_1mm=config.allow_unique_1mm,
                allow_reverse_complement=config.allow_reverse_complement,
            )
            sequence_matches[sample_id][spacer] = match
            sample_metrics[sample_id][f"{match.match_type.lower()}_reads"] += count
            run_metrics[f"{match.match_type.lower()}_reads"] += count
            if match.match_type == "EXACT":
                exact_counts[sample_id][match.guide_id] += count
            elif match.match_type == "UNIQUE_1MM":
                rescued_counts[sample_id][match.guide_id] += count

    spacer_rows = []
    spacer_calls = {}
    for sample_id in sample_names:
        extracted_total = sum(extracted_sequences[sample_id].values())
        metrics = sample_metrics[sample_id]
        anchor_matched_pairs = (
            metrics["extracted"]
            + metrics["spacer_length_invalid"]
            + metrics["spacer_invalid_bases"]
        )
        anchor_eligible = anchor_matched_pairs >= config.min_sample_anchor_reads
        ranked_sequences = {
            sequence: rank
            for rank, (sequence, _) in enumerate(
                sorted(
                    extracted_sequences[sample_id].items(),
                    key=lambda item: (-item[1], item[0]),
                ),
                start=1,
            )
        }
        for guide in guides:
            exact_reads = exact_counts[sample_id][guide.guide_id]
            rescued_reads = rescued_counts[sample_id][guide.guide_id]
            spacer_fraction = (
                exact_reads / anchor_matched_pairs if anchor_matched_pairs else 0.0
            )
            formally_present = (
                anchor_eligible
                and spacer_fraction >= config.min_guide_fraction
            )
            if formally_present:
                spacer_call = "PRESENT"
            elif exact_reads:
                spacer_call = "BACKGROUND"
            else:
                spacer_call = "ABSENT"
            spacer_calls[(sample_id, guide.guide_id)] = spacer_call
            exact_sequence_ranks = [
                ranked_sequences[sequence]
                for sequence, match in sequence_matches[sample_id].items()
                if match.match_type == "EXACT"
                and match.guide_id == guide.guide_id
            ]
            spacer_rows.append(
                {
                    "sample_id": sample_id,
                    "spacer_id": guide.guide_id,
                    "spacer_seq": guide.spacer_seq,
                    "exact_reads": exact_reads,
                    "unique_1mm_reads": rescued_reads,
                    "anchor_matched_pairs": anchor_matched_pairs,
                    "extracted_spacer_reads": extracted_total,
                    "spacer_fraction": spacer_fraction,
                    "spacer_rank": min(exact_sequence_ranks) if exact_sequence_ranks else "",
                    "call": spacer_call,
                    "sample_status": "",
                    "no_sgrna_reason": "",
                }
            )

    sample_rows = []
    sample_qc_rows = []
    sample_status_lookup = {}
    sample_reason_lookup = {}
    for sample_id in sample_names:
        extracted_total = sum(extracted_sequences[sample_id].values())
        barcode_matched_pairs = sample_metrics[sample_id]["barcode_matched_pairs"]
        metrics = sample_metrics[sample_id]
        anchor_matched_pairs = (
            metrics["extracted"]
            + metrics["spacer_length_invalid"]
            + metrics["spacer_invalid_bases"]
        )
        positive_spacer_rows = [
            row
            for row in spacer_rows
            if row["sample_id"] == sample_id and row["call"] == "PRESENT"
        ]
        if anchor_matched_pairs < config.min_sample_anchor_reads:
            sample_status = "NO_SGRNA"
            no_sgrna_reason = "LOW_ANCHOR_READS"
        elif positive_spacer_rows:
            sample_status = "SGRNA_DETECTED"
            no_sgrna_reason = ""
        else:
            sample_status = "NO_SGRNA"
            no_sgrna_reason = "NO_SPACER_ABOVE_THRESHOLD"
        sample_status_lookup[sample_id] = sample_status
        sample_reason_lookup[sample_id] = no_sgrna_reason
        sample_rows.append(
            {
                "sample_id": sample_id,
                "barcode_matched_pairs": barcode_matched_pairs,
                "anchor_matched_pairs": anchor_matched_pairs,
                "extracted_spacer_reads": extracted_total,
                "positive_spacer_count": len(positive_spacer_rows),
                "positive_spacer_ids": ";".join(
                    row["spacer_id"] for row in positive_spacer_rows
                ),
                "positive_spacer_reads": ";".join(
                    str(row["exact_reads"]) for row in positive_spacer_rows
                ),
                "positive_spacer_fractions": ";".join(
                    f"{row['spacer_fraction']:.2%}" for row in positive_spacer_rows
                ),
                "sample_status": sample_status,
                "no_sgrna_reason": no_sgrna_reason,
            }
        )
        sample_qc_rows.append(
            {
                "sample_id": sample_id,
                "barcode_matched_pairs": barcode_matched_pairs,
                "anchor_matched_pairs": anchor_matched_pairs,
                "extracted_spacer_reads": extracted_total,
                "exact_spacer_reads": metrics["exact_reads"],
                "unique_1mm_reads": metrics["unique_1mm_reads"],
                "ambiguous_reads": metrics["ambiguous_reads"],
                "unmatched_reads": metrics["unmatched_reads"],
                "anchor_match_rate": (
                    anchor_matched_pairs / barcode_matched_pairs
                    if barcode_matched_pairs else 0.0
                ),
                "spacer_extraction_rate": (
                    extracted_total / barcode_matched_pairs
                    if barcode_matched_pairs else 0.0
                ),
                "exact_spacer_rate": (
                    metrics["exact_reads"] / extracted_total if extracted_total else 0.0
                ),
                "sample_status": sample_status,
                "no_sgrna_reason": no_sgrna_reason,
            }
        )

    for row in spacer_rows:
        row["sample_status"] = sample_status_lookup[row["sample_id"]]
        row["no_sgrna_reason"] = sample_reason_lookup[row["sample_id"]]

    sample_detail_rows = {}
    for sample_id in sample_names:
        metrics = sample_metrics[sample_id]
        anchor_matched_pairs = (
            metrics["extracted"]
            + metrics["spacer_length_invalid"]
            + metrics["spacer_invalid_bases"]
        )
        rows = []
        for spacer, count in sorted(
            extracted_sequences[sample_id].items(),
            key=lambda item: (-item[1], item[0]),
        ):
            match = sequence_matches[sample_id][spacer]
            spacer_id = ""
            if match.guide_id is not None:
                spacer_id = match.guide_id
                spacer_call = (
                    spacer_calls[(sample_id, match.guide_id)]
                    if match.match_type == "EXACT"
                    else "AUXILIARY"
                )
            elif match.match_type == "AMBIGUOUS":
                spacer_id = ";".join(match.candidate_guide_ids)
                spacer_call = "AMBIGUOUS"
            else:
                spacer_call = "UNMATCHED"
            rows.append(
                {
                    "Sequence": spacer,
                    "Count": count,
                    "Percentage": (
                        f"{count / anchor_matched_pairs * 100:.2f}%"
                        if anchor_matched_pairs else "0.00%"
                    ),
                    "Match_Type": match.match_type,
                    "Spacer_ID": spacer_id,
                    "Spacer_Call": spacer_call,
                }
            )
        sample_detail_rows[sample_id] = rows

    for metric in (
        "total_pairs",
        "barcode_q30_fail_r1",
        "barcode_q30_fail_r2",
        "barcode_mismatch",
        "barcode_matched_pairs",
        "extracted",
        "exact_reads",
        "unique_1mm_reads",
        "ambiguous_reads",
        "unmatched_reads",
    ):
        run_metrics[metric] += 0
    anchor_matched_pairs = (
        run_metrics["extracted"]
        + run_metrics["spacer_length_invalid"]
        + run_metrics["spacer_invalid_bases"]
    )
    total_pairs = run_metrics["total_pairs"]
    barcode_matched_pairs = run_metrics["barcode_matched_pairs"]
    extracted_total = run_metrics["extracted"]
    run_metrics["anchor_matched_pairs"] = anchor_matched_pairs
    run_metrics["barcode_match_rate"] = (
        barcode_matched_pairs / total_pairs if total_pairs else 0.0
    )
    run_metrics["anchor_match_rate"] = (
        anchor_matched_pairs / barcode_matched_pairs if barcode_matched_pairs else 0.0
    )
    run_metrics["spacer_extraction_rate"] = (
        extracted_total / barcode_matched_pairs if barcode_matched_pairs else 0.0
    )
    run_metrics["exact_spacer_rate"] = (
        run_metrics["exact_reads"] / extracted_total if extracted_total else 0.0
    )
    return VectorTraceResult(
        sample_names=sample_names,
        guides=list(guides),
        spacer_rows=spacer_rows,
        sample_rows=sample_rows,
        sample_qc_rows=sample_qc_rows,
        sample_detail_rows=sample_detail_rows,
        run_metrics=dict(run_metrics),
    )


SPACER_OUTPUT_FIELDS = (
    "sample_id",
    "spacer_id",
    "spacer_seq",
    "exact_reads",
    "unique_1mm_reads",
    "anchor_matched_pairs",
    "extracted_spacer_reads",
    "spacer_fraction",
    "spacer_rank",
    "call",
    "sample_status",
    "no_sgrna_reason",
)
SAMPLE_OUTPUT_FIELDS = (
    "sample_id",
    "barcode_matched_pairs",
    "anchor_matched_pairs",
    "extracted_spacer_reads",
    "positive_spacer_count",
    "positive_spacer_ids",
    "positive_spacer_reads",
    "positive_spacer_fractions",
    "sample_status",
    "no_sgrna_reason",
)
SAMPLE_QC_FIELDS = (
    "sample_id",
    "barcode_matched_pairs",
    "anchor_matched_pairs",
    "extracted_spacer_reads",
    "exact_spacer_reads",
    "unique_1mm_reads",
    "ambiguous_reads",
    "unmatched_reads",
    "anchor_match_rate",
    "spacer_extraction_rate",
    "exact_spacer_rate",
    "sample_status",
    "no_sgrna_reason",
)
SAMPLE_DETAIL_FIELDS = (
    "Sequence",
    "Count",
    "Percentage",
    "Match_Type",
    "Spacer_ID",
    "Spacer_Call",
)


def _write_tsv(path, fieldnames, rows):
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
            delimiter="\t",
            extrasaction="ignore",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def _append_worksheet_table(worksheet, fieldnames, rows, start_row=1):
    for column_index, fieldname in enumerate(fieldnames, start=1):
        cell = worksheet.cell(start_row, column_index, fieldname)
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, start=start_row + 1):
        for column_index, fieldname in enumerate(fieldnames, start=1):
            worksheet.cell(row_index, column_index, row.get(fieldname, ""))
    worksheet.freeze_panes = worksheet.cell(start_row + 1, 1)
    for column_index, fieldname in enumerate(fieldnames, start=1):
        width = max(
            len(str(fieldname)),
            max((len(str(row.get(fieldname, ""))) for row in rows), default=0),
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(width + 2, 48)


def _write_vector_trace_workbook(path, result):
    if Workbook is None:
        raise RuntimeError("openpyxl is required to create vector_trace_summary.xlsx")
    workbook = Workbook()
    plate = workbook.active
    plate.title = "Plate_Result"
    plate.cell(1, 1, "Row")
    plate.cell(1, 1).font = Font(bold=True)
    for column in range(1, 13):
        plate.cell(1, column + 1, column).font = Font(bold=True)
    for row_index, row_letter in enumerate("ABCDEFGH", start=2):
        plate.cell(row_index, 1, row_letter).font = Font(bold=True)

    positive_by_sample = {sample: [] for sample in result.sample_names}
    for row in result.spacer_rows:
        if row["call"] == "PRESENT":
            positive_by_sample[row["sample_id"]].append(
                f"{row['spacer_id']} ({row['spacer_fraction']:.2%})"
            )
    sample_rows = {row["sample_id"]: row for row in result.sample_rows}
    for sample_id in result.sample_names:
        match = re.fullmatch(r"([A-Ha-h])(1[0-2]|[1-9])", sample_id.strip())
        if not match:
            continue
        row_index = ord(match.group(1).upper()) - ord("A") + 2
        column_index = int(match.group(2)) + 1
        values = positive_by_sample[sample_id]
        summary = sample_rows[sample_id]
        if summary["no_sgrna_reason"] == "LOW_ANCHOR_READS":
            empty_label = "NO_SGRNA [LOW_ANCHOR]"
        else:
            empty_label = "NO_SGRNA [NO_TARGET]"
        plate.cell(
            row_index,
            column_index,
            "\n".join(values) if values else empty_label,
        ).alignment = Alignment(wrap_text=True, vertical="top")
    plate.freeze_panes = "B2"
    plate.column_dimensions["A"].width = 7
    for column_index in range(2, 14):
        plate.column_dimensions[get_column_letter(column_index)].width = 24
    for row_index in range(2, 10):
        max_lines = max(
            len(str(plate.cell(row_index, column_index).value or "").splitlines())
            for column_index in range(2, 14)
        )
        plate.row_dimensions[row_index].height = min(max(24, max_lines * 15), 390)

    sample_sheet = workbook.create_sheet("Sample_Summary")
    _append_worksheet_table(sample_sheet, SAMPLE_OUTPUT_FIELDS, result.sample_rows)
    qc_sheet = workbook.create_sheet("QC")
    qc_sheet.append(["Run_Metric", "Count"])
    for cell in qc_sheet[1]:
        cell.font = Font(bold=True)
    for metric, count in sorted(result.run_metrics.items()):
        qc_sheet.append([metric, count])
    qc_sheet.append([])
    sample_start = qc_sheet.max_row + 1
    _append_worksheet_table(qc_sheet, SAMPLE_QC_FIELDS, result.sample_qc_rows, sample_start)
    workbook.save(path)


def _sample_detail_filename(sample_id):
    safe_sample_id = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", sample_id.strip())
    if not safe_sample_id:
        raise ValueError("sample_id cannot produce an empty detail filename")
    return f"{safe_sample_id}_stat.txt"


def _write_sample_detail_files(result, output_dir):
    os.makedirs(output_dir, exist_ok=False)
    used_filenames = set()
    for sample_id in result.sample_names:
        filename = _sample_detail_filename(sample_id)
        if filename in used_filenames:
            raise ValueError(
                f"sample IDs produce duplicate detail filename: {filename}"
            )
        used_filenames.add(filename)
        _write_tsv(
            os.path.join(output_dir, filename),
            SAMPLE_DETAIL_FIELDS,
            result.sample_detail_rows[sample_id],
        )


def write_vector_trace_outputs(result, output_dir, config, input_paths):
    if os.path.exists(output_dir) and os.listdir(output_dir):
        raise ValueError(f"output directory already exists and is not empty: {output_dir}")
    os.makedirs(output_dir, exist_ok=True)
    paths = {
        "vector_trace_summary.xlsx": os.path.join(
            output_dir,
            "vector_trace_summary.xlsx",
        ),
        "spacer_detection.tsv": os.path.join(output_dir, "spacer_detection.tsv"),
        "run_parameters.json": os.path.join(output_dir, "run_parameters.json"),
        "sample_details": os.path.join(output_dir, "sample_details"),
    }
    _write_tsv(
        paths["spacer_detection.tsv"],
        SPACER_OUTPUT_FIELDS,
        result.spacer_rows,
    )
    _write_vector_trace_workbook(paths["vector_trace_summary.xlsx"], result)
    _write_sample_detail_files(result, paths["sample_details"])
    normalized_inputs = {}
    for name, value in sorted(input_paths.items()):
        if isinstance(value, (list, tuple)):
            normalized_inputs[name] = [os.path.abspath(path) for path in value]
        else:
            normalized_inputs[name] = os.path.abspath(value)
    payload = {
        "mode": "vector-trace",
        "inputs": normalized_inputs,
        "parameters": asdict(config),
        "run_metrics": result.run_metrics,
        "formal_spacer_call_evidence": "EXACT_ONLY",
        "formal_spacer_fraction_denominator": "ANCHOR_MATCHED_READS",
        "one_mismatch_role": "AUXILIARY_ONLY",
    }
    with open(paths["run_parameters.json"], "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write("\n")
    return paths


def build_vector_trace_parser():
    parser = argparse.ArgumentParser(
        prog="hidogV11.py vector-trace",
        description=(
            "Detect anchor-adjacent sgRNA spacers by spacer ID and sequence "
            "without running genotyping."
        ),
    )
    parser.add_argument("-1", "--read1", "--r1", required=True, help="Raw R1 FASTQ(.gz)")
    parser.add_argument("-2", "--read2", "--r2", required=True, help="Raw R2 FASTQ(.gz)")
    parser.add_argument(
        "-b",
        "--barcode",
        required=True,
        help="Whitespace-delimited file: sample_id barcode_R1 barcode_R2",
    )
    reference_input = parser.add_mutually_exclusive_group(required=True)
    reference_input.add_argument(
        "-s",
        "--spacer-ref",
        nargs="+",
        help="One or more spacer FASTA/text files containing spacer IDs and sequences",
    )
    reference_input.add_argument(
        "-g",
        "--guide-manifest",
        help=(
            "Deprecated compatibility TSV: vector_id, guide_id, spacer_seq, "
            "gene_name, reference_group"
        ),
    )
    parser.add_argument("--anchor", required=True, help="Vector constant sequence adjacent to spacer")
    parser.add_argument(
        "--anchor-read",
        choices=("r1", "r2"),
        default="r1",
        help="Read containing anchor and spacer [default: r1]",
    )
    parser.add_argument(
        "--anchor-max-mismatches",
        type=int,
        default=0,
        help="Maximum anchor Hamming mismatches; tied best positions are rejected [default: 0]",
    )
    parser.add_argument(
        "--spacer-side",
        choices=("before", "after"),
        default="after",
        help="Spacer position relative to anchor [default: after]",
    )
    parser.add_argument(
        "--spacer-length",
        type=int,
        default=20,
        help="Expected spacer length [default: 20]",
    )
    parser.add_argument(
        "--barcode-spacer-length",
        type=int,
        default=4,
        help="Bases before each sample barcode [default: 4]",
    )
    parser.add_argument(
        "--barcode-length",
        type=int,
        default=4,
        help="Sample barcode length on each read [default: 4]",
    )
    parser.add_argument(
        "--barcode-q30",
        type=int,
        default=30,
        help="Minimum quality for every barcode base [default: 30]",
    )
    parser.add_argument(
        "--min-guide-fraction",
        type=float,
        default=0.05,
        help="Minimum exact fraction among anchor-matched reads [default: 0.05]",
    )
    parser.add_argument(
        "--min-guide-reads",
        type=int,
        default=10,
        help="Deprecated compatibility option; formal calls use fraction only",
    )
    parser.add_argument(
        "--min-sample-barcode-reads",
        type=int,
        default=2000,
        help=(
            "Deprecated compatibility option; sample eligibility uses anchor reads"
        ),
    )
    parser.add_argument(
        "--min-sample-anchor-reads",
        type=int,
        default=4000,
        help=(
            "Minimum anchor-matched read pairs required for an SGRNA_DETECTED "
            "sample call [default: 4000]"
        ),
    )
    parser.add_argument(
        "--allow-unique-1mm",
        dest="allow_unique_1mm",
        action="store_true",
        default=True,
        help="Count unique one-mismatch assignments as auxiliary evidence [default: enabled]",
    )
    parser.add_argument(
        "--no-unique-1mm",
        dest="allow_unique_1mm",
        action="store_false",
        help="Disable auxiliary unique one-mismatch assignment",
    )
    parser.add_argument(
        "--allow-reverse-complement",
        action="store_true",
        help="Also compare extracted spacers to reverse-complemented whitelist entries",
    )
    parser.add_argument(
        "--top-unmatched",
        type=int,
        default=20,
        help=(
            "Deprecated compatibility option; sample detail files always retain "
            "all unmatched sequences"
        ),
    )
    parser.add_argument(
        "-o",
        "--output",
        required=True,
        help="New or empty vector-trace output directory",
    )
    return parser


def parse_vector_trace_args(argv=None):
    return build_vector_trace_parser().parse_args(argv)


def vector_trace_main(argv=None):
    parser = build_vector_trace_parser()
    args = parser.parse_args(argv)
    for label, path in (
        ("read1", args.read1),
        ("read2", args.read2),
        ("barcode", args.barcode),
    ):
        if not os.path.isfile(path):
            parser.error(f"{label} file not found: {path}")
    if args.spacer_ref:
        for path in args.spacer_ref:
            if not os.path.isfile(path):
                parser.error(f"spacer reference file not found: {path}")
    elif not os.path.isfile(args.guide_manifest):
        parser.error(f"guide manifest file not found: {args.guide_manifest}")
    if os.path.exists(args.output) and os.listdir(args.output):
        parser.error(f"output directory already exists and is not empty: {args.output}")

    try:
        config = VectorTraceConfig(
            anchor=args.anchor,
            spacer_length=args.spacer_length,
            spacer_side=args.spacer_side,
            anchor_read=args.anchor_read,
            anchor_max_mismatches=args.anchor_max_mismatches,
            barcode_spacer_length=args.barcode_spacer_length,
            barcode_length=args.barcode_length,
            barcode_q30=args.barcode_q30,
            min_guide_fraction=args.min_guide_fraction,
            min_guide_reads=args.min_guide_reads,
            min_sample_barcode_reads=args.min_sample_barcode_reads,
            min_sample_anchor_reads=args.min_sample_anchor_reads,
            allow_unique_1mm=args.allow_unique_1mm,
            allow_reverse_complement=args.allow_reverse_complement,
            top_unmatched=args.top_unmatched,
        )
        if args.spacer_ref:
            spacer_records = load_spacer_references(
                args.spacer_ref,
                spacer_length=args.spacer_length,
            )
            guides = [
                GuideRecord(
                    vector_id=spacer.spacer_id,
                    guide_id=spacer.spacer_id,
                    spacer_seq=spacer.spacer_seq,
                    gene_name=spacer.spacer_id,
                    reference_group=spacer.spacer_id,
                )
                for spacer in spacer_records
            ]
            reference_inputs = {"spacer_references": args.spacer_ref}
        else:
            guides = load_guide_manifest(
                args.guide_manifest,
                spacer_length=args.spacer_length,
            )
            reference_inputs = {"guide_manifest": args.guide_manifest}
        result = analyze_vector_trace(
            read1=args.read1,
            read2=args.read2,
            barcode_file=args.barcode,
            guides=guides,
            config=config,
        )
        paths = write_vector_trace_outputs(
            result=result,
            output_dir=args.output,
            config=config,
            input_paths={
                "read1": args.read1,
                "read2": args.read2,
                "barcode": args.barcode,
                **reference_inputs,
            },
        )
    except (OSError, ValueError, RuntimeError) as exc:
        parser.error(str(exc))

    detected_samples = sum(
        row["sample_status"] == "SGRNA_DETECTED" for row in result.sample_rows
    )
    no_sgrna_samples = len(result.sample_rows) - detected_samples
    print(f"Vector trace completed for {len(result.sample_names)} sample(s).")
    print(f"Detected samples: {detected_samples}")
    print(f"No-sgRNA samples: {no_sgrna_samples}")
    print(f"Results: {os.path.abspath(args.output)}")
    print(f"Spacer detection: {os.path.abspath(paths['spacer_detection.tsv'])}")
    print(f"Sample details: {os.path.abspath(paths['sample_details'])}")
    return 0
